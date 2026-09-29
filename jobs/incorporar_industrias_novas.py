"""
incorporar_industrias_novas.py — Coloca na Base Mestre as indústrias que
estão na Receita (CNAE na Tabela DN da CNI) mas ainda não estavam na base.

Origem: data/processed/BASE_AMPLIADA_AL.csv, Tipo "Indústria fora da Base
Mestre" (gerado por jobs/carregar_base_ampliada.py). Em 07/2026 eram 317,
a maioria aberta em 2026 — a Base Mestre tinha sido gerada antes.

O que faz:
  1. acrescenta essas empresas ao BASE_MESTRE_COMERCIAL.csv, uma linha por
     empresa (matriz), com ORIGEM_REGISTRO = "Receita AAAA-MM (incorporada)".
     Colunas que dependem só do CNAE/porte (setor, subsetor, classificação
     SEBRAE...) são copiadas de empresas da própria base com o mesmo CNAE e
     porte — a mesma classificação que o BI_Project já deu a elas.
  2. tira essas empresas da base ampliada (não aparecem duas vezes)
  3. gera a lista para o comercial, com contatos, em
     data/contatos/INDUSTRIAS_NOVAS_<competencia>.xlsx (fora do Git)

Pode rodar de novo sem duplicar: quem já está na Base Mestre é ignorado.
Quando o BI_Project gerar uma Base Mestre nova, ela substitui esta.

Uso:
    python jobs/incorporar_industrias_novas.py --competencia 2026-07
"""

import argparse
import re
import sys
import unicodedata
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.tools.company_data import BASE_PATH  # noqa: E402
from src.tools.contatos import carregar_contatos  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
ARQ_AMPLIADA = RAIZ / "data" / "processed" / "BASE_AMPLIADA_AL.csv"
PASTA_LISTA = RAIZ / "data" / "contatos"
TIPO_LACUNA = "Indústria fora da Base Mestre"

# colunas que dependem só do CNAE (e do porte, no caso do SEBRAE)
COLS_POR_CNAE = ["SEBRAE_secao", "SEBRAE_setor", "SEBRAE_subsetor", "SEBRAE_e_industria", "SEBRAE_auditoria",
                 "SEBRAE_cnae_codigo_recuperado", "SEBRAE_cnae_divisao", "CNAE PRIMARIO"]
COLS_POR_CNAE_PORTE = ["SEBRAE_elegivel_sebrae", "SEBRAE_oportunidade_comercial", "SEBRAE_ELEGIVEL",
                       "SEBRAE_OPORTUNIDADE_ORIGINAL", "STATUS_SEBRAE", "SEBRAE_OPORTUNIDADE", "OPORTUNIDADE_SEBRAE"]
RELACIONAMENTO = {  # (TEM_SESI, TEM_SENAI) -> colunas de relacionamento, como na base
    (False, False): ("SEM RELACIONAMENTO", "PROSPECT", "NÃO", "SIM", "SIM"),
    (True, False): ("SOMENTE SESI", "CLIENTE", "SENAI", "NÃO", "SIM"),
    (False, True): ("SOMENTE SENAI", "CLIENTE", "SESI", "SIM", "NÃO"),
    (True, True): ("SESI + SENAI", "CLIENTE", "NÃO", "NÃO", "NÃO"),
}


def _norm(t) -> str:
    t = unicodedata.normalize("NFKD", str(t or ""))
    return "".join(c for c in t if not unicodedata.combining(c)).lower().strip()


def _cnae7(v) -> str:
    d = re.sub(r"\D", "", str(v or "").split(".")[0])
    return d.zfill(7) if d else ""


def _moda(df: pd.DataFrame, chave: list[str], cols: list[str]) -> pd.DataFrame:
    base = df.dropna(subset=chave)
    return base.groupby(chave)[cols].agg(lambda s: s.mode().iat[0] if not s.mode().empty else "")


def montar_linhas(novas: pd.DataFrame, base: pd.DataFrame, competencia: str) -> pd.DataFrame:
    base = base.copy()
    base["_cnae"] = base["SEBRAE_cnae_norm"].map(_cnae7)
    base["_div"] = base["_cnae"].str[:2]
    cols_cnae = [c for c in COLS_POR_CNAE if c in base]
    cols_porte = [c for c in COLS_POR_CNAE_PORTE if c in base]
    por_cnae = _moda(base, ["_cnae"], cols_cnae)
    por_div = _moda(base, ["_div"], cols_cnae)
    por_cnae_porte = _moda(base, ["_cnae", "Porte"], cols_porte)
    por_div_porte = _moda(base, ["_div", "Porte"], cols_porte)
    grafias = base["Municipio"].dropna().value_counts()           # grafia mais usada na base vence
    municipios = {}
    for m in grafias.index:
        municipios.setdefault(_norm(m), m)

    linhas = []
    for x in novas.to_dict("records"):
        cnae = _cnae7(x["cnae_principal"])
        div, porte = cnae[:2], x["Porte"]
        sesi = str(x.get("POSSUI_SESI", "")).upper() == "TRUE"
        senai = str(x.get("POSSUI_SENAI", "")).upper() == "TRUE"
        status, geral, cross, op_sesi, op_senai = RELACIONAMENTO[(sesi, senai)]
        cnpj = re.sub(r"\D", "", str(x["cnpj"])).zfill(14)
        municipio = municipios.get(_norm(x["Municipio"]), x["Municipio"])
        linha = {c: "" for c in base.columns if not c.startswith("_")}
        linha.update({
            "cnpj": cnpj, "CNPJ_NORMALIZADO": cnpj, "CNPJ_BASICO": cnpj[:8], "SEBRAE_cnpj": f"{cnpj}.0",
            "situacao_atual": "ATIVA", "SEBRAE_situacao_atual": "ATIVA",
            "razao_social": x["razao_social"], "SEBRAE_razao_social": x["razao_social"],
            "SETOR": "INDÚSTRIA", "SEBRAE_SETOR": "INDÚSTRIA", "EH_INDUSTRIA": "True",
            "Porte": porte, "SEBRAE_Porte": porte, "Municipio": municipio, "SEBRAE_Municipio": municipio,
            "SEBRAE_cnae_norm": f"{cnae}.0", "CNAE_DIVISAO": "0",
            "TEM_SESI": str(sesi), "TEM_SENAI": str(senai), "TEM_SESI_SENAI": str(sesi and senai),
            "STATUS_RELACIONAMENTO": status, "OPORTUNIDADE_GERAL": geral, "OPORTUNIDADE_CROSS_SELL": cross,
            "OPORTUNIDADE_SESI": op_sesi, "OPORTUNIDADE_SENAI": op_senai, "ENCONTRADO_SEBRAE": "True",
            "ORIGEM_REGISTRO": f"Receita {competencia} (incorporada)",
        })
        for tabela, chave, cols in ((por_cnae, cnae, cols_cnae), (por_div, div, cols_cnae)):
            if chave in tabela.index:
                for c in cols:
                    if not linha.get(c):
                        linha[c] = tabela.at[chave, c]
        for tabela, chave in ((por_cnae_porte, (cnae, porte)), (por_div_porte, (div, porte))):
            if chave in tabela.index:
                for c in cols_porte:
                    if not linha.get(c):
                        linha[c] = tabela.at[chave, c]
        if not linha.get("CNAE PRIMARIO"):
            linha["CNAE PRIMARIO"] = _norm(x.get("CNAE PRIMARIO", "")).capitalize()
        if not linha.get("SEBRAE_cnae_divisao") and div:
            linha["SEBRAE_cnae_divisao"] = f"{int(div)}.0"
        linhas.append(linha)
    return pd.DataFrame(linhas)


def lista_comercial(novas: pd.DataFrame, contatos: pd.DataFrame) -> pd.DataFrame:
    c = contatos.drop_duplicates("cnpj_basico").set_index("cnpj_basico") if not contatos.empty else pd.DataFrame()
    pega = lambda b, col: c.at[b, col] if (not c.empty and b in c.index and col in c.columns) else ""  # noqa: E731
    out = pd.DataFrame({
        "CNPJ": novas["cnpj"], "Razão social": novas["razao_social"], "Nome fantasia": novas.get("nome_fantasia", ""),
        "Município": novas["Municipio"], "Porte": novas["Porte"], "CNAE principal": novas["CNAE PRIMARIO"],
        "Início de atividade": pd.to_datetime(novas["data_inicio_atividade"], format="%Y%m%d", errors="coerce").dt.date,
        "SPE": novas["razao_social"].str.contains(r"\bSPE\b", case=False, na=False).map({True: "Sim", False: ""}),
        "Telefone": [pega(b, "telefone_1") for b in novas["cnpj_basico"]],
        "WhatsApp provável": [pega(b, "whatsapp_provavel") for b in novas["cnpj_basico"]],
        "E-mail": [pega(b, "email") for b in novas["cnpj_basico"]],
        "Responsável": [pega(b, "decisor") for b in novas["cnpj_basico"]],
    })
    return out.sort_values(["Início de atividade", "Razão social"], ascending=[False, True], na_position="last")


def salvar_excel(lista: pd.DataFrame, caminho: Path) -> None:
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(caminho, engine="openpyxl") as w:
        lista.to_excel(w, index=False, sheet_name="Indústrias novas")
        ws = w.sheets["Indústrias novas"]
        for j, col in enumerate(lista.columns, 1):
            cel = ws.cell(1, j)
            cel.font = Font(name="Arial", bold=True, color="FFFFFF")
            cel.fill = PatternFill("solid", fgColor="1F3864")
            largura = max(12, min(55, int(lista[col].astype(str).str.len().quantile(0.9)) + 2))
            ws.column_dimensions[get_column_letter(j)].width = largura
        for linha in ws.iter_rows(min_row=2):
            for cel in linha:
                cel.font = Font(name="Arial", size=10)
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--competencia", required=True, help="competência da Receita usada (ex.: 2026-07)")
    args = parser.parse_args()

    amp = pd.read_csv(ARQ_AMPLIADA, dtype=str, encoding="utf-8-sig").fillna("")
    base = pd.read_csv(BASE_PATH, dtype=str, low_memory=False, encoding="utf-8-sig")
    if "ORIGEM_REGISTRO" not in base:
        base["ORIGEM_REGISTRO"] = "Base Mestre (BI_Project)"
    ja = set(base["CNPJ_BASICO"].astype(str).str.zfill(8))
    novas = amp[(amp["Tipo"] == TIPO_LACUNA) & ~amp["cnpj_basico"].isin(ja)]
    if novas.empty:
        print("Nenhuma indústria nova para incorporar.")
        return

    linhas = montar_linhas(novas, base, args.competencia)
    base_nova = pd.concat([base, linhas[base.columns]], ignore_index=True)
    base_nova.to_csv(BASE_PATH, index=False, encoding="utf-8-sig")
    amp[amp["Tipo"] != TIPO_LACUNA].to_csv(ARQ_AMPLIADA, index=False, encoding="utf-8-sig")

    PASTA_LISTA.mkdir(parents=True, exist_ok=True)
    arq_lista = PASTA_LISTA / f"INDUSTRIAS_NOVAS_{args.competencia.replace('-', '')}.xlsx"
    salvar_excel(lista_comercial(novas, carregar_contatos()), arq_lista)

    print(f"Incorporadas à Base Mestre: {len(linhas):,} empresas")
    print(f"Base Mestre: {base['CNPJ_BASICO'].nunique():,} -> {base_nova['CNPJ_BASICO'].nunique():,} empresas (CNPJ raiz)")
    print(f"Lista para o comercial: {arq_lista}")


if __name__ == "__main__":
    main()
